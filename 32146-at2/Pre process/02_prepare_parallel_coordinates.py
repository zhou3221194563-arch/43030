#!/usr/bin/env python3
"""Create the long, normalized data sheet for chart 3: parallel coordinates."""

from pathlib import Path
from statistics import mean

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


INPUT_WORKBOOK = Path(__file__).resolve().parents[1] / "output" / "2026-Spring-32146-Ass2-V20_tableau_clean.xlsx"
OUTPUT_WORKBOOK = Path(__file__).resolve().parents[1] / "output" / "2026-Spring-32146-Ass2-V20_parallel_coordinates.xlsx"
SHEET = "Parallel_Coordinates"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def main():
    workbook = load_workbook(INPUT_WORKBOOK, read_only=True, data_only=True)
    source = workbook["Tableau_Data"]
    headers = [cell.value for cell in source[1]]
    records = [dict(zip(headers, row)) for row in source.iter_rows(min_row=2, values_only=True)]
    players = {}

    for record in records:
        minutes = number(record["Mins"])
        for player_name, role, seed in [
            (record["Champion"], "wins", record["Champion Seed"]),
            (record["Runner-up"], "runner_ups", record["Runner-up Seed"]),
        ]:
            player = players.setdefault(player_name, {"gender": record["Gender"], "wins": 0, "runner_ups": 0, "seeds": [], "mins": []})
            player[role] += 1
            if number(seed) is not None:
                player["seeds"].append(number(seed))
            if minutes is not None:
                player["mins"].append(minutes)

    summaries = []
    for name, player in players.items():
        if player["wins"] >= 5:
            finals = player["wins"] + player["runner_ups"]
            summaries.append({
                "Player": name,
                "Gender": player["gender"],
                "Champion Wins": player["wins"],
                "Runner-up Count": player["runner_ups"],
                "Win Rate": player["wins"] / finals,
                "Avg Seed": mean(player["seeds"]) if player["seeds"] else None,
                "Avg Mins": mean(player["mins"]) if player["mins"] else None,
            })
    summaries.sort(key=lambda item: (-item["Champion Wins"], item["Player"]))

    metrics = ["Champion Wins", "Runner-up Count", "Win Rate", "Avg Seed", "Avg Mins"]
    ranges = {metric: (min(item[metric] for item in summaries if item[metric] is not None), max(item[metric] for item in summaries if item[metric] is not None)) for metric in metrics}
    rows = []
    for summary in summaries:
        for metric in metrics:
            raw_value = summary[metric]
            minimum, maximum = ranges[metric]
            normalized = None if raw_value is None else 0 if minimum == maximum else (raw_value - minimum) / (maximum - minimum)
            rows.append([summary["Player"], summary["Gender"], metric, raw_value, normalized])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(["Player", "Gender", "Metric", "Raw Value", "Normalized Value"])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    for column, width in {"A": 20, "B": 12, "C": 22, "D": 14, "E": 18}.items():
        sheet.column_dimensions[column].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = "0.0"
        sheet.cell(row, 5).number_format = "0.000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(OUTPUT_WORKBOOK)


if __name__ == "__main__":
    main()
