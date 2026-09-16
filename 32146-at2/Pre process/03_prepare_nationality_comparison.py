#!/usr/bin/env python3
"""Create the country-and-result count sheet for chart 6."""

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


INPUT_WORKBOOK = Path(__file__).resolve().parents[1] / "output" / "2026-Spring-32146-Ass2-V20_tableau_clean.xlsx"
OUTPUT_WORKBOOK = Path(__file__).resolve().parents[1] / "output" / "2026-Spring-32146-Ass2-V20_nationality_comparison.xlsx"
SHEET = "Nationality_Comparison"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)


def main():
    workbook = load_workbook(INPUT_WORKBOOK, read_only=True, data_only=True)
    source = workbook["Tableau_Data"]
    headers = [cell.value for cell in source[1]]
    records = [dict(zip(headers, row)) for row in source.iter_rows(min_row=2, values_only=True)]
    champion_counts = Counter(record["Champion Country"] for record in records if record["Champion Country"])
    runner_up_counts = Counter(record["Runner-up Country"] for record in records if record["Runner-up Country"])
    countries = sorted(champion_counts.keys() | runner_up_counts.keys(), key=lambda country: (-(champion_counts[country] + runner_up_counts[country]), country))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(["Country", "Result", "Count"])
    for country in countries:
        sheet.append([country, "Champion", champion_counts[country]])
        sheet.append([country, "Runner-up", runner_up_counts[country]])
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    for column, width in {"A": 24, "B": 14, "C": 12}.items():
        sheet.column_dimensions[column].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 3).number_format = "#,##0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(OUTPUT_WORKBOOK)


if __name__ == "__main__":
    main()
