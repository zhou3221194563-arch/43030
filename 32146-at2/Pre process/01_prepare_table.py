#!/usr/bin/env python3
"""Build the clean Tableau_Data sheet from the supplied Wimbledon workbook."""

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "2026-Spring-32146-Ass2-V20.xlsx"
OUTPUT = ROOT / "output" / "2026-Spring-32146-Ass2-V20_tableau_clean.xlsx"
SOURCE_SHEET = "wimbledons_champions"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def completion_status(score):
    text = str(score or "").lower()
    if "walkover" in text or "w/o" in text:
        return "Walkover"
    if "retd" in text or "ret" in text:
        return "Retirement"
    return "Completed"


def main():
    source_workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    source = source_workbook[SOURCE_SHEET]
    rows = list(source.iter_rows(values_only=True))
    headers = list(rows[0])
    records = [dict(zip(headers, row)) for row in rows[1:]]
    champion_counts = Counter(record["Champion"] for record in records)

    output_workbook = Workbook()
    table = output_workbook.active
    table.title = "Tableau_Data"
    table_headers = headers + [
        "Total Sets Played",
        "Total Games Played",
        "Match Completion Status",
        "Duration Availability",
        "Champion Win Count",
        "Top Player (5+ titles)",
        "Decade",
        "Champion Seed Status",
        "Runner-up Seed Status",
    ]
    table.append(table_headers)

    set_prefixes = ["1st", "2nd", "3rd", "4th", "5th"]
    for record in records:
        set_pairs = [(number(record[f"{prefix}-won"]), number(record[f"{prefix}-loss"])) for prefix in set_prefixes]
        played_sets = [(won, lost) for won, lost in set_pairs if won is not None and lost is not None]
        total_games = sum(won + lost for won, lost in played_sets)
        year = record["Year"]
        champion_seed = number(record["Champion Seed"])
        runner_up_seed = number(record["Runner-up Seed"])
        table.append(
            [record[header] for header in headers]
            + [
                len(played_sets),
                total_games,
                completion_status(record["Score"]),
                "Recorded" if number(record["Mins"]) is not None else "Unavailable",
                champion_counts[record["Champion"]],
                "Yes" if champion_counts[record["Champion"]] >= 5 else "No",
                f"{year // 10 * 10}s",
                "Seeded" if champion_seed is not None else "Unseeded/unknown",
                "Seeded" if runner_up_seed is not None else "Unseeded/unknown",
            ]
        )

    for cell in table[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    table.freeze_panes = "A2"
    table.auto_filter.ref = table.dimensions
    output_workbook.save(OUTPUT)


if __name__ == "__main__":
    main()
